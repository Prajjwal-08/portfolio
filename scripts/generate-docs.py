import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_document():
    doc = docx.Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    def set_cell_background(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("Autonomous Git Repository & Database Lifecycle Guide")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(19, 20, 32)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(20)
    run_sub = sub_p.add_run("From Repository Clone to Autonomous Development, Database Seeding, Build Validation, and Automated Git Push")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 100, 110)

    # Section 1: Starting with Cloning
    h1 = doc.add_heading("1. Initial Phase: Cloning the Repository", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Every autonomous engineering task begins when the user provides a repository URL to clone or instructs the agent to clone the project. "
        "Depending on the authentication method chosen upfront, the clone command will differ:"
    )

    p_clone1 = doc.add_paragraph()
    p_clone1.paragraph_format.left_indent = Inches(0.25)
    r1 = p_clone1.add_run("Method A: Cloning via HTTPS with Personal Access Token (PAT) — Recommended\n")
    r1.bold = True
    p_clone1.add_run(
        "When using HTTPS without interactive prompts, clone directly with the PAT embedded or via credential injection:\n"
        "   git clone https://<YOUR_GITHUB_PAT>@github.com/<USERNAME>/<REPO_NAME>.git\n"
        "   cd <REPO_NAME>\n"
        "This embeds authentication for all subsequent fetch and push actions during the session."
    )

    p_clone2 = doc.add_paragraph()
    p_clone2.paragraph_format.left_indent = Inches(0.25)
    r2 = p_clone2.add_run("Method B: Cloning via SSH (Passphrase-less)\n")
    r2.bold = True
    p_clone2.add_run(
        "If an SSH key has already been linked to your GitHub account:\n"
        "   git clone git@github.com:<USERNAME>/<REPO_NAME>.git\n"
        "   cd <REPO_NAME>"
    )

    p_clone3 = doc.add_paragraph()
    p_clone3.paragraph_format.left_indent = Inches(0.25)
    r3 = p_clone3.add_run("Method C: Standard Public / Local Clone\n")
    r3.bold = True
    p_clone3.add_run(
        "If you already ran 'git clone' manually, the agent simply verifies the current working branch and remote origin:\n"
        "   git status\n"
        "   git remote -v"
    )

    # Section 2: Permissions
    h2 = doc.add_heading("2. Essential Permissions & Setup", level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)

    doc.add_heading("A. GitHub Permissions (Why Headless Push Hangs)", level=2)
    doc.add_paragraph(
        "AI agents execute commands inside automated background subshells. Windows Credential Manager GUI prompts or browser "
        "sign-in dialogs cannot appear. Without non-interactive credentials, commands like 'git push' hang indefinitely. "
        "To enable autonomous pushing:"
    )

    p_pat_steps = doc.add_paragraph()
    p_pat_steps.paragraph_format.left_indent = Inches(0.25)
    p_pat_steps.add_run(
        "1. Open GitHub Settings → Developer Settings → Personal access tokens → Fine-grained tokens.\n"
        "2. Choose 'Only select repositories' and select your target repository.\n"
        "3. Under 'Repository permissions', set 'Contents' to 'Read and write'. (CRITICAL: 'Read-only' fails with 403 Forbidden).\n"
        "4. Copy the generated token string (github_pat_...).\n"
        "5. The agent pushes seamlessly using: git push https://$env:GIT_TOKEN@github.com/<USER>/<REPO>.git <BRANCH>"
    )

    doc.add_heading("B. Cloud Database Permissions (Firebase Firestore)", level=2)
    doc.add_paragraph(
        "To allow autonomous seeding scripts and the frontend/admin panels to read and write without authentication roadblocks, "
        "configure your Firebase Console → Firestore Database → Rules as follows:"
    )

    code_p = doc.add_paragraph()
    code_p.paragraph_format.left_indent = Inches(0.3)
    r_code = code_p.add_run(
        "rules_version = '2';\n"
        "service cloud.firestore {\n"
        "  match /databases/{database}/documents {\n"
        "    match /todos/{document=**} { allow create: if true; allow read: if false; }\n"
        "    match /projects/{document=**} { allow read, write: if true; }\n"
        "    match /skills/{document=**} { allow read, write: if true; }\n"
        "    match /experience/{document=**} { allow read, write: if true; }\n"
        "    match /admin_users/{document=**} { allow read, write: if true; }\n"
        "    match /activity_logs/{document=**} { allow read, write: if true; }\n"
        "  }\n"
        "}"
    )
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(9.5)
    r_code.font.color.rgb = RGBColor(40, 40, 50)

    # Section 3: Step-by-Step Lifecycle
    h3 = doc.add_heading("3. Full Autonomous Lifecycle After Cloning", level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Once the repository is cloned and in place, the agent systematically executes the following workflow:"
    )

    phases = [
        ("Step 1: Environment & Architecture Audit", "Runs 'git status' and 'git branch'. Inspects file structure to identify framework type (Next.js App Router, Vite, Tailwind CSS, etc.)."),
        ("Step 2: Dependency Verification & Install", "Reviews package.json and executes 'npm install' non-interactively to ensure libraries (Firebase, Lucide, Framer Motion) are installed in node_modules."),
        ("Step 3: Database Connection & Automated Seeding", "Inspects static datasets (src/data/*.ts), creates standalone seed scripts (scripts/seed-db.ts), and executes via tsx to populate Firestore collections automatically."),
        ("Step 4: Feature Implementation & Fallback Layer", "Constructs data access abstractions (src/lib/db.ts) with offline fallbacks so the app never crashes if the DB is unreachable. Builds authenticated admin panel (/admin)."),
        ("Step 5: Automated Testing & Build Validation", "Runs 'npm run build' to perform full TypeScript type-checking and Next.js static page generation. Any build or crypto errors are automatically diagnosed and resolved."),
        ("Step 6: Git Staging, Committing & Pushing", "Stages files with 'git add', creates descriptive commit messages with 'git commit', and pushes to GitHub with 'git push' using non-interactive tokens."),
    ]

    for title, desc in phases:
        p_ph = doc.add_paragraph()
        p_ph.paragraph_format.left_indent = Inches(0.2)
        r_ph = p_ph.add_run(f"• {title}: ")
        r_ph.bold = True
        p_ph.add_run(desc)

    # Section 4: Troubleshooting
    h4 = doc.add_heading("4. Troubleshooting & Gotchas Matrix", level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    headers = ["Symptom / Error", "Root Cause", "Autonomous Fix"]
    widths = [Inches(2.0), Inches(2.2), Inches(2.3)]

    for idx, text in enumerate(headers):
        hdr_cells[idx].text = text
        hdr_cells[idx].width = widths[idx]
        set_cell_background(hdr_cells[idx], "131420")
        for p in hdr_cells[idx].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    data = [
        ("git push hangs indefinitely", "Windows Credential Manager is prompting for a GUI popup in a headless terminal.", "Authenticate using PAT in remote URL or passphrase-less SSH key."),
        ("fatal: 403 Forbidden on git push", "Fine-grained PAT was created with Read-only permissions on repository Contents.", "Update PAT permissions in GitHub to 'Contents: Read and write'."),
        ("PERMISSION_DENIED on Firestore", "Firestore rules default to locking all writes ('allow write: if false;').", "Publish matching rules in Firebase Console allowing target collections."),
        ("crypto.subtle is undefined", "Web Cryptography API is only exposed in HTTPS or localhost secure contexts.", "Implement pure JavaScript SHA-256 fallback function."),
        ("Unstyled HTML / Broken layout", "Corrupted or stale Next.js cache in .next folder.", "Delete .next directory, rebuild/restart dev server, and hard refresh browser (Ctrl+F5)."),
    ]

    for row_data in data:
        row_cells = table.add_row().cells
        for idx, text in enumerate(row_data):
            row_cells[idx].text = text
            row_cells[idx].width = widths[idx]
            for p in row_cells[idx].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9.5)

    filename = "Autonomous_Git_Repository_Management_Guide_v2.docx"
    doc.save(filename)
    print(f"Word document generated successfully as {filename}!")

if __name__ == "__main__":
    create_document()
